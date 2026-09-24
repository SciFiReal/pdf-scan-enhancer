import logging
import re
from pathlib import Path

logger = logging.getLogger(__name__)

SUPPORTED_FORMATS = {
    "docx": "Word 文档（.docx）",
    "epub": "电子书（.epub）",
    "html": "网页（.html）",
    "txt": "纯文本（.txt）",
}


class ExportEngine:
    """导出引擎：将 Markdown 转换为 DOCX / EPUB / HTML / TXT"""

    @staticmethod
    def export(md_path: str, output_format: str, output_path: str = "") -> str:
        md_file = Path(md_path)
        if not md_file.exists():
            raise FileNotFoundError(f"Markdown 文件不存在：{md_path}")

        content = md_file.read_text(encoding="utf-8")
        fmt = output_format.lower().strip(".")

        if not output_path:
            output_path = str(md_file.parent / f"{md_file.stem}.{fmt}")

        if fmt == "docx":
            ExportEngine._to_docx(content, output_path, md_file.stem)
        elif fmt == "epub":
            ExportEngine._to_epub(content, output_path, md_file.stem)
        elif fmt == "html":
            ExportEngine._to_html(content, output_path, md_file.stem)
        elif fmt == "txt":
            ExportEngine._to_txt(content, output_path)
        else:
            raise ValueError(f"不支持的导出格式：{output_format}")

        logger.info("导出完成：%s", output_path)
        return output_path

    @staticmethod
    def merge_pdfs(pdf_paths: list[str], output_path: str) -> str:
        """合并多个 PDF 文件为一个"""
        import pikepdf

        if not pdf_paths:
            raise ValueError("没有可合并的 PDF 文件")

        if len(pdf_paths) == 1:
            import shutil
            shutil.copy2(pdf_paths[0], output_path)
            return output_path

        merged = pikepdf.Pdf.new()
        for pdf_path in pdf_paths:
            src = pikepdf.open(pdf_path)
            merged.pages.extend(src.pages)

        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        merged.save(output_path)
        merged.close()

        logger.info("PDF 合并完成：%s（%d 个文件）", output_path, len(pdf_paths))
        return output_path

    @staticmethod
    def export_multiple_md_to_epub(
        md_paths: list[str], output_path: str, title: str = "",
    ) -> str:
        """将多个 Markdown 文件合并为一本 EPUB，每个文件作为一个章节"""
        from ebooklib import epub

        book = epub.EpubBook()
        book.set_identifier("pdf-scan-enhancer-merged")
        book.set_title(title or "Document")
        book.set_language("zh-CN")
        book.add_author("PDF Scan Enhancer")

        css = epub.EpubItem(
            uid="style",
            file_name="style/default.css",
            media_type="text/css",
            content="""
body { font-family: serif; line-height: 1.8; padding: 1em; }
h1, h2, h3 { margin-top: 1.5em; }
p { text-align: justify; margin: 0.8em 0; }
em { color: #555; }
table { border-collapse: collapse; width: 100%; }
th, td { border: 1px solid #ccc; padding: 6px 10px; }
""".encode("utf-8"),
        )
        book.add_item(css)

        chapters = []
        for i, md_path in enumerate(md_paths):
            md_file = Path(md_path)
            content = md_file.read_text(encoding="utf-8")
            chapter_title = md_file.stem

            first_line = content.strip().split("\n")[0].strip() if content.strip() else ""
            if first_line.startswith("# "):
                chapter_title = first_line.lstrip("# ").strip()

            try:
                import markdown
                html_content = markdown.markdown(
                    content, extensions=["tables", "fenced_code", "toc"],
                )
            except ImportError:
                html_content = ExportEngine._md_to_html_basic(content)

            c = epub.EpubHtml(
                title=chapter_title,
                file_name=f"chap_{i+1:02d}.xhtml",
                lang="zh-CN",
            )
            c.content = f"<html><body>{html_content}</body></html>"
            c.add_item(css)
            book.add_item(c)
            chapters.append(c)

        book.toc = [(c,) for c in chapters]
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        book.spine = ["nav"] + chapters

        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        epub.write_epub(output_path, book, {})

        logger.info("EPUB 合并导出完成：%s（%d 章）", output_path, len(chapters))
        return output_path

    @staticmethod
    def _to_txt(content: str, output_path: str):
        text = re.sub(r"^#{1,6}\s+", "", content, flags=re.MULTILINE)
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        text = re.sub(r"\*(.+?)\*", r"\1", text)
        text = re.sub(r"!\[.*?\]\(.*?\)", "", text)
        text = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", text)
        text = re.sub(r"^---+$", "", text, flags=re.MULTILINE)

        Path(output_path).write_text(text.strip(), encoding="utf-8")

    @staticmethod
    def _to_html(content: str, output_path: str, title: str = ""):
        try:
            import markdown
            body = markdown.markdown(
                content,
                extensions=["tables", "fenced_code", "toc"],
            )
        except ImportError:
            body = ExportEngine._md_to_html_basic(content)

        html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title or "Document"}</title>
<style>
  body {{
    max-width: 800px;
    margin: 40px auto;
    padding: 0 20px;
    font-family: "Noto Serif SC", "Source Han Serif CN", "SimSun", serif;
    font-size: 16px;
    line-height: 1.8;
    color: #333;
  }}
  h1, h2, h3, h4, h5, h6 {{
    margin-top: 1.5em;
    margin-bottom: 0.5em;
    font-weight: 600;
  }}
  h1 {{ font-size: 1.8em; border-bottom: 2px solid #eee; padding-bottom: 0.3em; }}
  h2 {{ font-size: 1.4em; }}
  p {{ margin: 0.8em 0; text-align: justify; }}
  em {{ font-style: italic; color: #555; }}
  table {{ border-collapse: collapse; width: 100%; margin: 1em 0; }}
  th, td {{ border: 1px solid #ddd; padding: 8px 12px; text-align: left; }}
  th {{ background: #f5f5f5; font-weight: 600; }}
  hr {{ border: none; border-top: 1px solid #ddd; margin: 2em 0; }}
  blockquote {{ border-left: 4px solid #ddd; margin: 1em 0; padding: 0.5em 1em; color: #666; }}
  code {{ background: #f5f5f5; padding: 2px 6px; border-radius: 3px; font-size: 0.9em; }}
  pre {{ background: #f5f5f5; padding: 1em; border-radius: 4px; overflow-x: auto; }}
</style>
</head>
<body>
{body}
</body>
</html>"""
        Path(output_path).write_text(html, encoding="utf-8")

    @staticmethod
    def _md_to_html_basic(content: str) -> str:
        lines = content.split("\n")
        html_lines = []
        in_paragraph = False

        for line in lines:
            stripped = line.strip()

            if not stripped:
                if in_paragraph:
                    html_lines.append("</p>")
                    in_paragraph = False
                continue

            if stripped.startswith("# "):
                html_lines.append(f"<h1>{stripped[2:]}</h1>")
            elif stripped.startswith("## "):
                html_lines.append(f"<h2>{stripped[3:]}</h2>")
            elif stripped.startswith("### "):
                html_lines.append(f"<h3>{stripped[4:]}</h3>")
            elif stripped.startswith("---"):
                html_lines.append("<hr>")
            elif stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**"):
                html_lines.append(f"<p><em>{stripped.strip('*')}</em></p>")
            elif stripped.startswith("- "):
                html_lines.append(f"<li>{stripped[2:]}</li>")
            else:
                if not in_paragraph:
                    html_lines.append("<p>")
                    in_paragraph = True
                html_lines.append(stripped)

        if in_paragraph:
            html_lines.append("</p>")

        return "\n".join(html_lines)

    @staticmethod
    def _to_docx(content: str, output_path: str, title: str = ""):
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        style = doc.styles["Normal"]
        style.font.size = Pt(11)
        style.font.name = "SimSun"
        style.paragraph_format.line_spacing = 1.5

        paragraphs = content.split("\n\n")

        for para_text in paragraphs:
            lines = para_text.strip().split("\n")
            if not lines:
                continue

            first_line = lines[0].strip()

            if first_line.startswith("# "):
                p = doc.add_heading(first_line[2:], level=1)
            elif first_line.startswith("## "):
                p = doc.add_heading(first_line[3:], level=2)
            elif first_line.startswith("### "):
                p = doc.add_heading(first_line[4:], level=3)
            elif first_line.startswith("---"):
                p = doc.add_paragraph()
                p.add_run("─" * 40)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                continue
            elif first_line.startswith("|") and "|" in first_line[1:]:
                ExportEngine._add_table_to_docx(doc, para_text)
                continue
            elif first_line.startswith("- "):
                for line in lines:
                    line = line.strip()
                    if line.startswith("- "):
                        doc.add_paragraph(line[2:], style="List Bullet")
                continue
            else:
                full_text = "\n".join(l.strip() for l in lines)
                is_italic = full_text.startswith("*") and full_text.endswith("*") and not full_text.startswith("**")
                if is_italic:
                    full_text = full_text.strip("*")
                    p = doc.add_paragraph()
                    run = p.add_run(full_text)
                    run.italic = True
                else:
                    clean = re.sub(r"\*\*(.+?)\*\*", r"\1", full_text)
                    clean = re.sub(r"\*(.+?)\*", r"\1", clean)
                    clean = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", clean)
                    doc.add_paragraph(clean)

        doc.save(output_path)

    @staticmethod
    def _add_table_to_docx(doc, table_text: str):
        rows = []
        for line in table_text.strip().split("\n"):
            line = line.strip()
            if line.startswith("|") and not re.match(r"^\|[\s\-:|]+\|$", line):
                cells = [c.strip() for c in line.strip("|").split("|")]
                rows.append(cells)

        if not rows:
            return

        num_cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=num_cols, style="Table Grid")

        for i, row_data in enumerate(rows):
            for j, cell_text in enumerate(row_data):
                if j < num_cols:
                    table.cell(i, j).text = cell_text

    @staticmethod
    def _to_epub(content: str, output_path: str, title: str = ""):
        from ebooklib import epub

        book = epub.EpubBook()
        book.set_identifier("pdf-scan-enhancer-export")
        book.set_title(title or "Document")
        book.set_language("zh-CN")
        book.add_author("PDF Scan Enhancer")

        try:
            import markdown
            html_content = markdown.markdown(
                content, extensions=["tables", "fenced_code", "toc"],
            )
        except ImportError:
            html_content = ExportEngine._md_to_html_basic(content)

        css = epub.EpubItem(
            uid="style",
            file_name="style/default.css",
            media_type="text/css",
            content="""
body { font-family: serif; line-height: 1.8; padding: 1em; }
h1, h2, h3 { margin-top: 1.5em; }
p { text-align: justify; margin: 0.8em 0; }
em { color: #555; }
table { border-collapse: collapse; width: 100%; }
th, td { border: 1px solid #ccc; padding: 6px 10px; }
""".encode("utf-8"),
        )
        book.add_item(css)

        chapters = []
        sections = content.split("\n\n## ")

        if len(sections) <= 1:
            c1 = epub.EpubHtml(
                title=title or "全文",
                file_name="chap_01.xhtml",
                lang="zh-CN",
            )
            c1.content = f"<html><body>{html_content}</body></html>"
            c1.add_item(css)
            chapters.append(c1)
        else:
            for i, section in enumerate(sections):
                if i == 0 and section.strip().startswith("# "):
                    lines = section.strip().split("\n")
                    sec_title = lines[0].lstrip("# ").strip()
                elif i == 0:
                    sec_title = "前言"
                else:
                    sec_title = section.split("\n")[0].strip()
                    section = "\n".join(section.split("\n")[1:])

                try:
                    import markdown as md_lib
                    sec_html = md_lib.markdown(
                        section, extensions=["tables", "fenced_code"],
                    )
                except ImportError:
                    sec_html = ExportEngine._md_to_html_basic(section)

                c = epub.EpubHtml(
                    title=sec_title,
                    file_name=f"chap_{i+1:02d}.xhtml",
                    lang="zh-CN",
                )
                c.content = f"<html><body><h1>{sec_title}</h1>{sec_html}</body></html>"
                c.add_item(css)
                chapters.append(c)

        for c in chapters:
            book.add_item(c)

        book.toc = [(c,) for c in chapters]
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        book.spine = ["nav"] + chapters

        epub.write_epub(output_path, book, {})
